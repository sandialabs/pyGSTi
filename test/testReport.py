import unittest
import warnings
import pickle
import pygsti
from pygsti.construction import std1Q_XYI as std

import numpy as np
import os

class ReportTestCase(unittest.TestCase):

    def setUp(self):
        #Set GateSet objects to "strict" mode for testing
        pygsti.objects.GateSet._strict = True

        self.targetGateset = std.gs_target
        datagen_gateset = self.targetGateset.depolarize(gate_noise=0.05, spam_noise=0.1)
        
        self.fiducials = std.fiducials
        self.germs = std.germs
        self.specs = pygsti.construction.build_spam_specs(self.fiducials, effect_labels=['E0']) #only use the first EVec

        self.gateLabels = self.targetGateset.gates.keys() # also == std.gates
        self.lgstStrings = pygsti.construction.list_lgst_gatestrings(self.specs, self.gateLabels)

        self.maxLengthList = [0,1,2,4,8]
        
        self.lsgstStrings = pygsti.construction.make_lsgst_lists(
            self.gateLabels, self.fiducials, self.fiducials, self.germs, self.maxLengthList)

        self.ds = pygsti.objects.DataSet(fileToLoadFrom="cmp_chk_files/reportgen.dataset")

        # RUN BELOW LINES TO GENERATE ANALYSIS DATASET
        #ds = pygsti.construction.generate_fake_data(datagen_gateset, lsgstStrings[-1], nSamples=1000,
        #                                            sampleError='binomial', seed=100)
        #ds.save("cmp_chk_files/reportgen.dataset")

        gs_lgst = pygsti.do_lgst(self.ds, self.specs, self.targetGateset, svdTruncateTo=4, verbosity=0)
        gs_lgst_go = pygsti.optimize_gauge(gs_lgst,"target",targetGateset=self.targetGateset)
        self.gs_clgst = pygsti.contract(gs_lgst_go, "CPTP")
        self.gs_clgst_tp = pygsti.contract(self.gs_clgst, "vSPAM")
        self.gs_clgst_tp.set_all_parameterizations("TP")

        try:
            import pptx
            self.have_python_pptx = True
        except ImportError:
            warnings.warn("**** IMPORT: Cannot import pptx (python-pptx), and so" +
                         " Powerpoint slide generation tests have been disabled.")
            self.have_python_pptx = False


        #Compute results for MC2GST
        lsgst_gatesets_prego = pygsti.do_iterative_mc2gst(self.ds, self.gs_clgst, self.lsgstStrings, verbosity=0,
                                                          minProbClipForWeighting=1e-6, probClipInterval=(-1e6,1e6),
                                                          returnAll=True)
        lsgst_gatesets = [ pygsti.optimize_gauge(gs, "target", targetGateset=self.targetGateset,
                                                 gateWeight=1,spamWeight=0.001) for gs in lsgst_gatesets_prego]

        self.results = pygsti.report.Results()
        self.results.init_Ls_and_germs("chi2", self.targetGateset, self.ds, self.gs_clgst,
                                       self.maxLengthList, self.germs,
                                       lsgst_gatesets, self.lsgstStrings, self.fiducials, self.fiducials, 
                                       pygsti.construction.repeat_with_max_length, False, None, lsgst_gatesets_prego)
        self.results.parameters.update({'minProbClip': 1e-6, 'minProbClipForWeighting': 1e-4,
                                        'probClipInterval': (-1e6,1e6), 'radius': 1e-4,
                                        'weights': None, 'defaultDirectory': "temp_test_files",
                                        'defaultBasename': "MyDefaultReportName" } )


        #Compute results for MLGST with TP constraint
        lsgst_gatesets_TP = pygsti.do_iterative_mlgst(self.ds, self.gs_clgst_tp, self.lsgstStrings, verbosity=0,
                                                   minProbClip=1e-4, probClipInterval=(-1e6,1e6),
                                                   returnAll=True)
        lsgst_gatesets_TP = [ pygsti.optimize_gauge(gs, "target", targetGateset=self.targetGateset, constrainToTP=True,
                                                 gateWeight=1,spamWeight=0.001) for gs in lsgst_gatesets_TP]

        self.results_logL = pygsti.report.Results()
        self.results_logL.init_Ls_and_germs("logl", self.targetGateset, self.ds, self.gs_clgst_tp, self.maxLengthList, self.germs,
                                     lsgst_gatesets_TP, self.lsgstStrings, self.fiducials, self.fiducials, 
                                     pygsti.construction.repeat_with_max_length, True)


            
        

    def checkFile(self, fn):
        linesToTest = open("temp_test_files/%s" % fn).readlines()
        linesOK = open("cmp_chk_files/%s" % fn).readlines()
        self.assertEqual(linesToTest,linesOK)



class TestReport(ReportTestCase):
    
    def test_reports_chi2_noCIs(self):

        self.results.create_full_report_pdf(filename="temp_test_files/full_reportA.pdf", confidenceLevel=None,
                                         debugAidsAppendix=False, gaugeOptAppendix=False,
                                         pixelPlotAppendix=False, whackamoleAppendix=False)
        self.results.create_brief_report_pdf(filename="temp_test_files/brief_reportA.pdf", confidenceLevel=None)
        self.results.create_presentation_pdf(filename="temp_test_files/slidesA.pdf", confidenceLevel=None,
                                           debugAidsAppendix=False, pixelPlotAppendix=False, whackamoleAppendix=False)
        if self.have_python_pptx:
            self.results.create_presentation_ppt(filename="temp_test_files/slidesA.ppt", confidenceLevel=None,
                                                 debugAidsAppendix=False, pixelPlotAppendix=False, whackamoleAppendix=False)

        #Run again using default filenames
        self.results.create_full_report_pdf(filename="auto", confidenceLevel=None,
                                         debugAidsAppendix=False, gaugeOptAppendix=False,
                                         pixelPlotAppendix=False, whackamoleAppendix=False)
        self.results.create_brief_report_pdf(filename="auto", confidenceLevel=None)
        self.results.create_presentation_pdf(filename="auto", confidenceLevel=None,
                                           debugAidsAppendix=False, pixelPlotAppendix=False, whackamoleAppendix=False)
        if self.have_python_pptx:
            self.results.create_presentation_ppt(filename="auto", confidenceLevel=None,
                                                 debugAidsAppendix=False, pixelPlotAppendix=False, whackamoleAppendix=False)

        #Compare the text files, assume if these match the PDFs are equivalent
        self.checkFile("full_reportA.tex")
        self.checkFile("brief_reportA.tex")
        self.checkFile("slidesA.tex")



    def test_reports_chi2_wCIs(self):

        self.results.create_full_report_pdf(filename="temp_test_files/full_reportB.pdf", confidenceLevel=95,
                                         debugAidsAppendix=True, gaugeOptAppendix=True,
                                         pixelPlotAppendix=True, whackamoleAppendix=True,
                                         verbosity=2)
        self.results.create_full_report_pdf(filename="temp_test_files/full_reportB-noGOpt.pdf", confidenceLevel=95,
                                         debugAidsAppendix=True, gaugeOptAppendix=False,
                                         pixelPlotAppendix=True, whackamoleAppendix=True) # to test blank GOpt tables
        self.results.create_brief_report_pdf(filename="temp_test_files/brief_reportB.pdf", confidenceLevel=95, verbosity=2)
        self.results.create_presentation_pdf(filename="temp_test_files/slidesB.pdf", confidenceLevel=95,
                                           debugAidsAppendix=True, pixelPlotAppendix=True, whackamoleAppendix=True,
                                           verbosity=2)
        if self.have_python_pptx:
            self.results.create_presentation_ppt(filename="temp_test_files/slidesB.ppt", confidenceLevel=95,
                                                 debugAidsAppendix=True, pixelPlotAppendix=True, whackamoleAppendix=True,
                                                 verbosity=2)

        #Compare the text files, assume if these match the PDFs are equivalent
        self.checkFile("full_reportB.tex")
        self.checkFile("full_reportB_appendices.tex")
        self.checkFile("brief_reportB.tex")
        self.checkFile("slidesB.tex")


    def test_reports_chi2_nonMarkCIs(self):
        #Non-markovian error bars (negative confidenceLevel) & tooltips
        self.results.create_full_report_pdf(filename="temp_test_files/full_reportE.pdf", confidenceLevel=-95,
                                         debugAidsAppendix=True, gaugeOptAppendix=True,
                                         pixelPlotAppendix=True, whackamoleAppendix=True,
                                         verbosity=2, tips=True)
        self.results.create_brief_report_pdf(filename="temp_test_files/brief_reportE.pdf", confidenceLevel=-95,
                                             verbosity=2, tips=True)

        #Compare the text files, assume if these match the PDFs are equivalent
        self.checkFile("full_reportE.tex")
        self.checkFile("full_reportE_appendices.tex")


    def test_reports_logL_TP_noCIs(self):

        self.results_logL.create_full_report_pdf(filename="temp_test_files/full_reportC.pdf", confidenceLevel=None,
                                                 debugAidsAppendix=False, gaugeOptAppendix=False,
                                                 pixelPlotAppendix=False, whackamoleAppendix=False,
                                                 verbosity=2)
        self.results_logL.create_brief_report_pdf(filename="temp_test_files/brief_reportC.pdf", confidenceLevel=None, verbosity=2)
        self.results_logL.create_presentation_pdf(filename="temp_test_files/slidesC.pdf", confidenceLevel=None,
                                                  debugAidsAppendix=False, pixelPlotAppendix=False, whackamoleAppendix=False,
                                                  verbosity=2)
        self.results_logL.create_general_report_pdf(filename="temp_test_files/general_reportC.pdf", confidenceLevel=None,
                                                    verbosity=2)

        if self.have_python_pptx:
            self.results_logL.create_presentation_ppt(filename="temp_test_files/slidesC.ppt", confidenceLevel=None,
                                                      debugAidsAppendix=False, pixelPlotAppendix=False, whackamoleAppendix=False,
                                                      verbosity=2)

        ##Compare the text files, assume if these match the PDFs are equivalent
        self.checkFile("full_reportC.tex")
        self.checkFile("brief_reportC.tex")
        self.checkFile("slidesC.tex")
        #self.checkFile("general_reportC.tex")


    def test_reports_logL_TP_wCIs(self):

        self.results_logL.create_full_report_pdf(filename="temp_test_files/full_reportD.pdf", confidenceLevel=95,
                                                 debugAidsAppendix=True, gaugeOptAppendix=True,
                                                 pixelPlotAppendix=True, whackamoleAppendix=True,
                                                 verbosity=2)
        self.results_logL.create_brief_report_pdf(filename="temp_test_files/brief_reportD.pdf", confidenceLevel=95, verbosity=2)
        self.results_logL.create_presentation_pdf(filename="temp_test_files/slidesD.pdf", confidenceLevel=95,
                                                  debugAidsAppendix=True, pixelPlotAppendix=True, whackamoleAppendix=True,
                                                  verbosity=2)
        self.results_logL.create_general_report_pdf(filename="temp_test_files/general_reportD.pdf", confidenceLevel=95,
                                                    verbosity=2, tips=True) #test tips here too

        if self.have_python_pptx:
            self.results_logL.create_presentation_ppt(filename="temp_test_files/slidesD.ppt", confidenceLevel=95,
                                                      debugAidsAppendix=True, pixelPlotAppendix=True, whackamoleAppendix=True,
                                                      verbosity=2)

        ##Compare the text files, assume if these match the PDFs are equivalent
        self.checkFile("full_reportD.tex")
        self.checkFile("full_reportD_appendices.tex")
        self.checkFile("brief_reportD.tex")
        self.checkFile("slidesD.tex")
        #self.checkFile("general_reportD.tex")



    def test_extra_logL_TP_figs_and_tables(self):

        #Run a few tests to generate tables & figures we don't use in reports
        self.results_logL.tables["chi2ProgressTable"]
        self.results_logL.tables["logLProgressTable"]
        self.results_logL.figures["bestEstimateSummedColorBoxPlot"]
        self.results_logL.figures["blankBoxPlot"]
        self.results_logL.figures["blankSummedBoxPlot"]
        self.results_logL.figures["directLGSTColorBoxPlot"]
        self.results_logL.figures["directLGSTDeviationColorBoxPlot"]
        with self.assertRaises(KeyError):
            self.results_logL.figures["FooBar"]
        with self.assertRaises(KeyError):
            self.results_logL._specials['FooBar']

        #Run tests to generate tables we don't use in reports
        self.results_logL.tables["bestGatesetVsTargetAnglesTable"]



    def test_table_generation(self):
        import pygsti.report.generation as gen
        formats = ['latex','html','py','ppt'] #all the formats we know
        tableclass = "cssClass"
        longtable = False
        confidenceRegionInfo = None

        gateset = pygsti.io.load_gateset("cmp_chk_files/analysis.gateset")
        ds = pygsti.objects.DataSet(fileToLoadFrom="cmp_chk_files/analysis.dataset")

        chi2, chi2Hessian = pygsti.chi2(ds, gateset, returnHessian=True)
        ci = pygsti.obj.ConfidenceRegion(gateset, chi2Hessian, 95.0,
                                         hessianProjection="std")
        gateset_tp = pygsti.contract(gateset,"TP"); gateset_tp.set_all_parameterizations("TP")
        chi2, chi2Hessian_TP = pygsti.chi2(ds, gateset_tp, returnHessian=True)
        ci_TP = pygsti.obj.ConfidenceRegion(gateset_tp, chi2Hessian_TP, 95.0,
                                            hessianProjection="std")

        chi2, chi2Hessian_tgt = pygsti.chi2(ds, std.gs_target, returnHessian=True)
        ci_tgt = pygsti.obj.ConfidenceRegion(std.gs_target, chi2Hessian_tgt, 95.0,
                                         hessianProjection="std")
        target_tp = std.gs_target.copy(); target_tp.set_all_parameterizations("TP")
        chi2, chi2Hessian_tgt_TP = pygsti.chi2(ds, target_tp, returnHessian=True)
        ci_TP_tgt = pygsti.obj.ConfidenceRegion(target_tp, chi2Hessian_tgt_TP, 95.0,
                                            hessianProjection="std")

        gateset_2q = pygsti.construction.build_gateset( 
            [4], [('Q0','Q1')],['GIX','GIY','GXI','GYI','GCNOT'], 
            [ "I(Q0):X(pi/2,Q1)", "I(Q0):Y(pi/2,Q1)", "X(pi/2,Q0):I(Q1)", "Y(pi/2,Q0):I(Q1)", "CX(pi,Q0,Q1)" ],
            prepLabels=['rho0'], prepExpressions=["0"], effectLabels=['E0','E1','E2'], effectExpressions=["0","1","2"], 
            spamdefs={'upup': ('rho0','E0'), 'updn': ('rho0','E1'), 'dnup': ('rho0','E2'),
                           'dndn': ('rho0','remainder') }, basis="pp" )
        

        #tests which fill in the cracks of the full-report tests
        tab = gen.get_gateset_spam_table(gateset, None)
        tab_wCI = gen.get_gateset_spam_table(gateset, ci)
        gen.get_gateset_spam_table(gateset, None)
        gen.get_gateset_gates_table(gateset_tp, ci_TP) #test zero-padding
        gen.get_unitary_gateset_gates_table(std.gs_target, ci_tgt) #unitary gates w/CIs
        gen.get_unitary_gateset_gates_table(target_tp, ci_TP_tgt) #unitary gates w/CIs
        gen.get_gateset_closest_unitary_table(gateset_2q, None) #test higher-dim gateset
        gen.get_gateset_closest_unitary_table(gateset, ci) #test with CIs (long...)
        gen.get_gateset_rotn_axis_table(std.gs_target, None) #try to get "--"s and "X"s to display
        gen.get_chi2_progress_table([0], [gateset_tp], [ [('Gx',)],], ds) #TP case
        gen.get_chi2_confidence_region(gateset_tp, ds, 95) #TP case

        gen.get_gatestring_multi_table([ [('Gx',),('Gz',)], [('Gy',)] ],
                                       ["list1","list2"], commonTitle=None) #commonTitle == None case w/diff length lists

        with self.assertRaises(ValueError):
            gen.get_unitary_gateset_gates_table(std.gs_target, ci) #gateset-CI mismatch
        with self.assertRaises(ValueError):
            gen.get_gateset_spam_parameters_table(std.gs_target, ci) #gateset-CI mismatch

        #Test ReportTable object
        rowLabels = tab.keys()
        row1Data = tab[rowLabels[0]]
        colLabels = row1Data.keys()

        self.assertTrue(rowLabels, tab.row_names)
        self.assertTrue(colLabels, tab.col_names)
        self.assertTrue(len(rowLabels), tab.num_rows)
        self.assertTrue(len(colLabels), tab.num_cols)

        el00 = tab[rowLabels[0]][colLabels[0]]
        self.assertTrue( rowLabels[0] in tab )
        self.assertTrue( tab.has_key(rowLabels[0]) )

        table_len = len(tab)
        self.assertEqual(table_len, tab.num_rows)

        table_as_str = str(tab)
        table_wCI_as_str = str(tab_wCI)
        row1a = tab.row(key=rowLabels[0])
        col1a = tab.col(key=colLabels[0])
        row1b = tab.row(index=0)
        col1b = tab.col(index=0)
        self.assertEqual(row1a,row1b)
        self.assertEqual(col1a,col1b)

        with self.assertRaises(KeyError):
            tab['foobar']
        with self.assertRaises(KeyError):
            tab.row(key='foobar') #invalid key
        with self.assertRaises(ValueError):
            tab.row(index=100000) #out of bounds
        with self.assertRaises(ValueError):
            tab.row() #must specify key or index
        with self.assertRaises(ValueError):
            tab.row(key='foobar',index=1) #cannot specify key and index
        with self.assertRaises(KeyError):
            tab.col(key='foobar') #invalid key
        with self.assertRaises(ValueError):
            tab.col(index=100000) #out of bounds
        with self.assertRaises(ValueError):
            tab.col() #must specify key or index
        with self.assertRaises(ValueError):
            tab.col(key='foobar',index=1) #cannot specify key and index

        with self.assertRaises(ValueError):
            tab.render(fmt="foobar") #invalid format            



        #LogL case tests
        gen.get_logl_progress_table([0], [gateset_tp], [ [('Gx',)],], ds) # logL case
        gen.get_logl_progress_table([0], [gateset], [ [('Gx',)],], ds) # logL case

        gen.get_logl_confidence_region(gateset_tp, ds, 95,
                                       gatestring_list=None, probClipInterval=(-1e6,1e6),
                                       minProbClip=1e-4, radius=1e-4, hessianProjection="std")
        gen.get_logl_confidence_region(gateset, ds, 95,
                                       gatestring_list=None, probClipInterval=(-1e6,1e6),
                                       minProbClip=1e-4, radius=1e-4, hessianProjection="std")

        
    def test_table_formatting(self):
        vec = np.array( [1.0,2.0,3.0] )
        mx = np.identity( 2, 'd' )
        rank3Tensor = np.zeros( (2,2,2), 'd')
        f = 10.0
        l = [10.0, 20.0]
        
        class weirdType:
            def __init__(self):
                pass
            def __str__(self):
                return "weird"
        w = weirdType()

        print "Float formatting"
        pygsti.report.html.html(f)
        pygsti.report.latex.latex(f)
        pygsti.report.ppt.ppt(f)

        print "List formatting"
        pygsti.report.html.html(l)
        pygsti.report.latex.latex(l)
        pygsti.report.ppt.ppt(l)

        print "Arbitrary class formatting"
        pygsti.report.html.html(w)
        pygsti.report.latex.latex(w)
        pygsti.report.ppt.ppt(w)

        print "Vector formatting"
        pygsti.report.html.html(vec)
        pygsti.report.latex.latex(vec)
        pygsti.report.ppt.ppt(vec)

        print "Vector formatting (w/brackets)"
        pygsti.report.html.html(vec, brackets=True)
        pygsti.report.latex.latex(vec, brackets=True)
        pygsti.report.ppt.ppt(vec, brackets=True)

        print "Matrix formatting"
        pygsti.report.html.html_matrix(mx, fontsize=8, brackets=False)
        pygsti.report.latex.latex_matrix(mx, fontsize=8, brackets=False)
        pygsti.report.ppt.ppt_matrix(mx, fontsize=8, brackets=False)

        print "Value formatting"
        ROUND = 2; complxAsPolar=True
        for complxAsPolar in (True,False):
            for x in (0.001,0.01,1.0,10.0,100.0,1000.0,10000.0,1.0+1.0j,10j,1.0+1e-10j,1e-10j,"N/A"):
                pygsti.report.html.html_value(x, ROUND, complxAsPolar)
                pygsti.report.latex.latex_value(x, ROUND, complxAsPolar)
                pygsti.report.ppt.ppt_value(x, ROUND, complxAsPolar)        

        with self.assertRaises(ValueError):
            pygsti.report.html.html(rank3Tensor)

        with self.assertRaises(ValueError):
            pygsti.report.latex.latex(rank3Tensor)

        with self.assertRaises(ValueError):
            pygsti.report.ppt.ppt(rank3Tensor)


    def test_plotting(self):
        test_data = np.array( [[1e-8,1e-7,1e-6,1e-5],
                               [1e-4,1e-3,1e-2,1e-1],
                               [1.0,10.0,100.0,1000.0],
                               [1.0e4,1.0e5,1.0e6,1.0e7]],'d' )
        cmap = pygsti.report.plotting.StdColormapFactory('seq', n_boxes=10, vmin=0, vmax=1, dof=1)
        gstFig = pygsti.report.plotting.color_boxplot( test_data, cmap, size=(10,10), prec="compacthp",save_to="temp_test_files/test.pdf")
        gstFig = pygsti.report.plotting.color_boxplot( test_data, cmap, size=(10,10), prec="compact",save_to="")
        gstFig = pygsti.report.plotting.color_boxplot( test_data, cmap, size=(10,10), prec=3,save_to="")
        gstFig = pygsti.report.plotting.color_boxplot( test_data, cmap, size=(10,10), prec=-3,save_to="")
        gstFig = pygsti.report.plotting.color_boxplot( test_data, cmap, size=(10,10), prec="foobar",colorbar=True,save_to="")
        gstFig.check()

        pygsti.report.plotting.gateset_with_lgst_gatestring_estimates( [('Gx','Gx')], self.ds, self.specs,
                                                                       self.targetGateset,includeTargetGates=False,
                                                                       gateStringLabels=None, svdTruncateTo=4, verbosity=0)

        gsWithGxgx = pygsti.report.plotting.focused_mc2gst_gatesets( 
            pygsti.construction.gatestring_list([('Gx','Gx')]), self.ds, self.specs, self.gs_clgst)

    def test_reportables(self):
        #Test that None is returned when qty cannot be computed
        qty = pygsti.report.reportables.compute_dataset_qty("FooBar",self.ds)
        self.assertIsNone(qty)
        qty = pygsti.report.reportables.compute_gateset_qty("FooBar",self.gs_clgst)
        self.assertIsNone(qty)
        qty = pygsti.report.reportables.compute_gateset_dataset_qty("FooBar",self.gs_clgst, self.ds)
        self.assertIsNone(qty)
        qty = pygsti.report.reportables.compute_gateset_gateset_qty("FooBar",self.gs_clgst, self.gs_clgst)
        self.assertIsNone(qty)

        #test ignoring gate strings not in dataset        
        qty = pygsti.report.reportables.compute_dataset_qty("gate string length", self.ds,
                                                            pygsti.construction.gatestring_list([('Gx','Gx'),('Gfoobar',)]) )
        qty = pygsti.report.reportables.compute_gateset_dataset_qty("prob(plus) diff", self.gs_clgst, self.ds,
                                                            pygsti.construction.gatestring_list([('Gx','Gx'),('Gfoobar',)]) )
        qty_str = str(qty) #test __str__

        #Test gateset gates mismatch
        from pygsti.construction import std1Q_XY as stdXY          
        with self.assertRaises(ValueError):
            qty = pygsti.report.reportables.compute_gateset_gateset_qty(
                "Gx fidelity",std.gs_target, stdXY.gs_target) #Gi missing from 2nd gateset
        with self.assertRaises(ValueError):
            qty = pygsti.report.reportables.compute_gateset_gateset_qty(
                "Gx fidelity",stdXY.gs_target, std.gs_target) #Gi missing from 1st gateset



    def test_results_object(self):
        results = pygsti.report.Results()
        results.init_single("logl", self.targetGateset, self.ds, self.gs_clgst,
                            self.lgstStrings, False, self.targetGateset)

        results.parameters.update(
            {'minProbClip': 1e-6, 'minProbClipForWeighting': 1e-4,
             'probClipInterval': (-1e6,1e6), 'radius': 1e-4,
             'weights': None, 'defaultDirectory': "temp_test_files",
             'defaultBasename': "MyDefaultReportName" } )
        
        results.create_full_report_pdf(
            filename="temp_test_files/singleReport.pdf")
        results.create_brief_report_pdf(
            filename="temp_test_files/singleBrief.pdf")
        results.create_presentation_pdf(
            filename="temp_test_files/singleSlides.pdf")
        if self.have_python_pptx:
            results.create_presentation_ppt(
                filename="temp_test_files/singleSlides.ppt", pptTables=True)

        #test tree splitting of hessian
        results.parameters['memLimit'] = 10*(1024)**2 #10MB
        results.create_brief_report_pdf(confidenceLevel=95,
            filename="temp_test_files/singleBriefMemLimit.pdf")
        results.parameters['memLimit'] = 10 #10 bytes => too small
        with self.assertRaises(MemoryError):
            results.create_brief_report_pdf(confidenceLevel=90,
               filename="temp_test_files/singleBriefMemLimit.pdf")


        #similar test for chi2 hessian
        results2 = pygsti.report.Results()
        results2.init_single("chi2", self.targetGateset, self.ds, self.gs_clgst,
                            self.lgstStrings, False, self.targetGateset)
        results2.parameters.update(
            {'minProbClip': 1e-6, 'minProbClipForWeighting': 1e-4,
             'probClipInterval': (-1e6,1e6), 'radius': 1e-4,
             'weights': None, 'defaultDirectory': "temp_test_files",
             'defaultBasename': "MyDefaultReportName" } )
        results2.parameters['memLimit'] = 10*(1024)**2 #10MB
        results2.create_brief_report_pdf(confidenceLevel=95,
            filename="temp_test_files/singleBriefMemLimit2.pdf")
        results2.parameters['memLimit'] = 10 #10 bytes => too small
        with self.assertRaises(MemoryError):
            results2.create_brief_report_pdf(confidenceLevel=90,
               filename="temp_test_files/singleBriefMemLimit2.pdf")




        results_str = str(results)
        tableNames = results.tables.keys()
        figNames = results.figures.keys()
        for g in results.gatesets:
            s = str(g)
        for g in results.gatestring_lists:
            s = str(g)
        s = str(results.dataset)
        s = str(results.options)

        self.assertTrue(results.tables.has_key(tableNames[0]))

        with self.assertRaises(KeyError):
            x = results.tables.get('foobar')
        with self.assertRaises(ValueError):
            results.tables['newKey'] = "notAllowed"
        with self.assertRaises(NotImplementedError):
            for x in results.tables: # cannot __iter__
                print x
        with self.assertRaises(NotImplementedError):
            for x in results.tables.iteritems(): # cannot iter
                print x
        with self.assertRaises(NotImplementedError):
            for x in results.tables.values(): # cannot iter
                print x

        pkl = pickle.dumps(results)
        results_copy = pickle.loads(pkl)
        self.assertEqual(tableNames, results_copy.tables.keys())
        self.assertEqual(figNames, results_copy.figures.keys())
        #self.assertEqual(results.options, results_copy.options) #need to add equal test to ResultsOptions
        self.assertEqual(results.parameters, results_copy.parameters)

        results2 = pygsti.report.Results()
        results2.options.template_path = "/some/path/to/templates"
        results2.options.latex_cmd = "myCustomLatex"

        #bad objective function name
        results_badObjective = pygsti.report.Results()
        #results_badObjective.init_single("foobar", self.targetGateset, self.ds, self.gs_clgst,
        #                                 self.lgstStrings, False)
        results_badObjective.init_Ls_and_germs("foobar", self.targetGateset, self.ds, self.gs_clgst, [0], self.germs,
                                               [self.gs_clgst], [self.lgstStrings], self.fiducials, self.fiducials, 
                                               pygsti.construction.repeat_with_max_length, True)
        
        with self.assertRaises(ValueError):
            results_badObjective._get_confidence_region(95) 
        with self.assertRaises(ValueError):
            results_badObjective._specials['DirectLongSeqGatesets']
        with self.assertRaises(ValueError):
            results_badObjective.create_full_report_pdf(filename="temp_test_files/badReport.pdf")
        with self.assertRaises(ValueError):
            results_badObjective.create_presentation_pdf(filename="temp_test_files/badSlides.pdf")
        if self.have_python_pptx:
            with self.assertRaises(ValueError):
                results_badObjective.create_presentation_ppt(filename="temp_test_files/badSlides.pptx")
        

    def test_report_figure_object(self):
        axes = {'dummy': "matplotlib axes"}
        fig = pygsti.report.figure.ReportFigure(axes, {})
        fig.set_extra_info("extra!")
        self.assertEqual(fig.get_extra_info(), "extra!")
        
        with self.assertRaises(ValueError):
            fig.pickledAxes = "not-a-pickle-string" #corrupt pickled string so get unpickling error
            fig.save_to("temp_test_files/test.figure")
    

        
if __name__ == "__main__":
    unittest.main(verbosity=2)
